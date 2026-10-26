"""Vuelca el Excel nuevo del plan rodilla a docs/plan-rodilla/volcado.json y
compara la hoja de volumen con la versión anterior (la del commit c3ece19, que
es la que la app tenía empotrada antes del plan rodilla).

    python tools/volcado_excel.py

volcado.json: {hoja: {celda: {"v": valor calculado, "f": fórmula}}} con todas
las celdas con valor de "Volumen (1910-1112)" y "Seguimiento rodilla".
diff_volumen.json / diff_volumen.txt: celdas de la hoja Volumen que cambian.
"""
import datetime
import io
import json
import pathlib
import subprocess
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs" / "plan-rodilla"
NUEVO = DOCS / "Plan_8sem_vol_definitivo_con_cardio.xlsx"
# La versión anterior se lee de git: el fichero de la raíz pasa a ser el nuevo.
ANTERIOR = "c3ece19:Plan_8sem_vol_definitivo_con_cardio.xlsx"
HOJAS = ["Volumen (1910-1112)", "Seguimiento rodilla"]


def js(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


def volcar(path, hojas):
    f = openpyxl.load_workbook(path)
    v = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for h in hojas:
        if h not in f.sheetnames:
            continue
        celdas = {}
        for row in f[h].iter_rows():
            for c in row:
                if c.value is None:
                    continue
                d = {"v": js(v[h][c.coordinate].value)}
                if isinstance(c.value, str) and c.value.startswith("="):
                    d["f"] = c.value
                elif not isinstance(c.value, str) and not isinstance(c.value, (int, float)):
                    d["f"] = str(c.value)   # fórmulas de array, etc.
                celdas[c.coordinate] = d
        out[h] = celdas
    return out


def filas_por_ejercicio(celdas):
    """{(sesión, ejercicio): {col: texto}} para las tablas de sesión y el bloque
    de cardio. Las filas se desplazan entre versiones; así se comparan por
    contenido y no por posición."""
    def t(ref):
        c = celdas.get(ref)
        return None if c is None else c.get("f", c.get("v"))
    filas = {}
    sesion = None
    ultima = max(int("".join(ch for ch in r if ch.isdigit())) for r in celdas)
    for r in range(1, ultima + 1):
        a = t(f"A{r}")
        if isinstance(a, str) and a.startswith("(") and ")" in a[:4]:
            sesion = a[1:a.index(")")]
            continue
        if isinstance(a, str) and a.startswith("BLOQUE DE CARDIO"):
            sesion = "CARDIO"
            continue
        if a in (None, "Ejercicio", "Día / Sesión") or sesion is None:
            if a is None and sesion != "CARDIO":
                sesion = None
            continue
        if isinstance(a, str) and (a.startswith("Series semanales") or a.startswith("REGLAS")):
            sesion = None
            continue
        clave = (sesion, a.strip() if sesion != "CARDIO" else f"{a.strip()} | {t(f'B{r}')}")
        filas[clave] = {col: t(f"{col}{r}") for col in "BCDEFGHIJK"} | {"fila": r}
    return filas


def main():
    nuevo = volcar(NUEVO, HOJAS)
    (DOCS / "volcado.json").write_text(json.dumps(nuevo, ensure_ascii=False, indent=1), encoding="utf-8")
    viejo_bytes = subprocess.run(["git", "show", ANTERIOR], cwd=ROOT, capture_output=True, check=True).stdout
    viejo = volcar(io.BytesIO(viejo_bytes), HOJAS[:1])
    h = HOJAS[0]
    a, b = viejo[h], nuevo[h]
    diff = []
    for ref in sorted(set(a) | set(b), key=lambda r: (openpyxl.utils.cell.coordinate_to_tuple(r))):
        va = a.get(ref, {}).get("f", a.get(ref, {}).get("v"))
        vb = b.get(ref, {}).get("f", b.get(ref, {}).get("v"))
        if va != vb:
            diff.append({"celda": ref, "antes": va, "ahora": vb})
    (DOCS / "diff_volumen.json").write_text(json.dumps(diff, ensure_ascii=False, indent=1), encoding="utf-8")
    with open(DOCS / "diff_volumen.txt", "w", encoding="utf-8") as fh:
        for d in diff:
            fh.write(f"{d['celda']}\n  antes: {d['antes']!r}\n  ahora: {d['ahora']!r}\n")
    fa, fb = filas_por_ejercicio(a), filas_por_ejercicio(b)
    with open(DOCS / "diff_por_ejercicio.txt", "w", encoding="utf-8") as fh:
        for k in sorted(set(fa) | set(fb), key=lambda k: (str(k[0]), fb.get(k, fa.get(k))["fila"])):
            if k not in fa:
                fh.write(f"ALTA {k[0]} · {k[1]!r} (fila {fb[k]['fila']})\n")
                for col in "BCDEFGHIJK":
                    fh.write(f"    {col}: {fb[k][col]!r}\n")
            elif k not in fb:
                fh.write(f"BAJA/RENOMBRE {k[0]} · {k[1]!r} (fila antigua {fa[k]['fila']})\n")
            else:
                cambios = [col for col in "BCDEFGHIJK" if fa[k][col] != fb[k][col]]
                if cambios:
                    fh.write(f"CAMBIA {k[0]} · {k[1]!r} (fila {fa[k]['fila']} -> {fb[k]['fila']})\n")
                    for col in cambios:
                        fh.write(f"    {col}: {fa[k][col]!r} -> {fb[k][col]!r}\n")
    print(f"volcado: {sum(len(c) for c in nuevo.values())} celdas · diff Volumen: {len(diff)} celdas")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
