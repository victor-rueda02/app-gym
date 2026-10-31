"""Simula una edición futura del plan TOCANDO SOLO EL EXCEL (en una copia, en tools/out/) y
comprueba que la app la sigue sin cambiar código:

    python tools/prueba_edicion.py

Ediciones: un ejercicio nuevo en mitad de la hoja (desplaza todo lo de debajo), un
renombrado que conserva el id (y por tanto el historial), una prescripción, un nivel de
carrera y una regla más. Después regenera las fórmulas (tools/excel_app.py), construye el
plan con el mismo código que la app (Node) y comprueba el resultado, las celdas
combinadas y la tabla de series por grupo (con el paquete `formulas`, si está).
"""
import base64
import json
import pathlib
import shutil
import subprocess
import sys

import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from excel_app import ROOT, VOL, XLSX, insertar_filas, quitar_filas, regenerar  # noqa: E402

COPIA = ROOT / "tools" / "out" / "prueba_edicion.xlsx"

PLAN_JS = r"""
const fs = require('fs');
const bytes = fs.readFileSync(process.argv[1]);
const src = ['vendor/fflate.js', 'js/lector-xlsx.js', 'js/plan.js'].map((f) => fs.readFileSync(f, 'utf8')).join('\n;\n');
const PLAN = new Function('module', 'exports', 'originalXlsxBytes', src + ';return PLAN;')(undefined, undefined, () => new Uint8Array(bytes));
const ej = Object.values(PLAN.BLOCK_VOL).flatMap((s) => s.exercises.map((x) => ({ sesion: s.id, id: x.id, name: x.name, target: x.target, labels: x.pres.map((p) => p.label) })));
process.stdout.write(JSON.stringify({ ej, carrera: PLAN.CARRERA_NIVELES, reglas: PLAN.VOL_BLOCK_RULES,
  series: [1, 2, 3, 4, 5, 6, 7, 8].map((n) => PLAN.seriesPorGrupo(n)) }));
"""


# El motor de progresión con un Excel cualquiera: S2 y S3 en verde (mañanas a 2, cada
# sesión con «durante» bajo y la carrera con su nivel previo) y la entrada superada.
MOTOR_JS = r"""
const fs = require('fs');
const bytes = fs.readFileSync(process.argv[1]);
const src = ['vendor/fflate.js', 'js/lector-xlsx.js', 'js/plan.js', 'js/rodilla-data.js', 'js/rodilla-calc.js']
  .map((f) => fs.readFileSync(f, 'utf8')).join('\n;\n');
const { PLAN, RODILLA_DATA, RodillaCalc: C } = new Function('module', 'exports', 'originalXlsxBytes',
  src + ';return { PLAN, RODILLA_DATA, RodillaCalc };')(undefined, undefined, () => new Uint8Array(bytes));
const L = PLAN.ESCALERAS.simples[0];
const dias = {};
for (const d of RODILLA_DATA.dias.filter((x) => x.fecha < '2026-11-12')) {
  dias[d.fecha] = { decline: 2 };
  if ('LJV'.includes(d.dia)) dias[d.fecha].during = 2;   // días de gimnasio (pierna 2: jueves en S2, viernes desde S3)
  if ('XD'.includes(d.dia)) Object.assign(dias[d.fecha], { during: 1, back60: 'Sí' });
}
delete dias['2026-10-26'].decline;   // el lunes de S2 sin medir, como en la realidad
const entrada = { S2: { walkPain: 'Sí', walkBack60: 'Sí', tdPain: 'Sí', hopsPain: 'Sí' } };
const prog = (d) => C.progresion(C.tablaDias(RODILLA_DATA.dias, d), PLAN, {}, entrada, '2026-11-12');
const p = prog(dias);
const ej = Object.values(PLAN.BLOCK_VOL).flatMap((s) => s.exercises).find((x) => x.prog === 'extension');
// Con el viernes de S3 en Mantener (la mañana del domingo, a 48 h, peor) no sube en S4.
const p2 = prog({ ...dias, '2026-11-08': { decline: 3 } });
process.stdout.write(JSON.stringify({
  escalera: { max: L.max, niveles: L.niveles },
  niveles: ['S2', 'S3', 'S4'].map((s) => p.semanas[s].simples.extension.nivel),
  S4: C.prescripcion(ej, p.semanas.S4, PLAN, 4).label,
  S4retro: C.prescripcion(ej, p2.semanas.S4, PLAN, 4).label,
}));
"""


def plan_de(path):
    out = subprocess.run(["node", "-e", PLAN_JS, str(path)], cwd=ROOT, capture_output=True, encoding="utf-8")
    if out.returncode:
        raise SystemExit(out.stderr)
    return json.loads(out.stdout)


def fila(ws, texto, col="A"):
    return next(r for r in range(1, ws.max_row + 1) if str(ws[f"{col}{r}"].value or "").startswith(texto))


def main():
    antes = plan_de(XLSX)
    COPIA.parent.mkdir(exist_ok=True)
    shutil.copyfile(XLSX, COPIA)
    wb = openpyxl.load_workbook(COPIA, rich_text=True)
    ws = wb[VOL]

    # 1. Ejercicio nuevo al final de la tabla del sábado (antes de «Autorregulación»).
    auto = fila(ws, "Autorregulación")
    insertar_filas(ws, auto, 1)
    for col, v in zip("ABCDEFGHIJKL", ["Plancha lateral", "Abs", "-"] + ["2 @ RIR1/0"] * 7
                      + ["Nueva: 2 series desde S2.", "id=plancha-lateral"]):
        ws[f"{col}{auto}"] = v
    # 2. Renombrado: el id (columna App) se queda, así que el historial no se pierde.
    r = fila(ws, "Elev laterales mancuerna")
    ws[f"A{r}"] = "Elevaciones laterales con mancuerna"
    # 3. Una prescripción.
    ws[f"J{fila(ws, 'Elev laterales polea')}"] = "4 @ RIR1/0"
    # 4. Un nivel de carrera.
    rc = next(r for r in range(1, ws.max_row + 1) if "escalera=carrera" in str(ws[f"L{r}"].value or "")
              and "domingo" not in str(ws[f"L{r}"].value))
    ws[f"H{rc}"] = "N5 · 3×(7'/1')"
    # 5. Una regla más al final de las reglas del bloque.
    ultima = fila(ws, "Fases hacia el fútbol")
    insertar_filas(ws, ultima + 1, 1, estilo_de=ultima)
    ws[f"A{ultima + 1}"] = "Regla de prueba [LÓGICA]."
    wb.save(COPIA)
    regenerar(COPIA)
    despues = plan_de(COPIA)

    errores = []
    def check(cond, msg):
        if not cond:
            errores.append(msg)

    nuevo = [e for e in despues["ej"] if e["id"] == "plancha-lateral"]
    check(len(nuevo) == 1 and nuevo[0]["sesion"] == "VOL-S" and nuevo[0]["labels"][1] == "2 @ RIR1/0",
          f"ejercicio nuevo: {nuevo}")
    check(len(despues["ej"]) == len(antes["ej"]) + 1, "nº de ejercicios")
    ren = [e for e in despues["ej"] if e["id"] == "elev-laterales-mancuerna"]
    check(len(ren) == 1 and ren[0]["name"] == "Elevaciones laterales con mancuerna", f"renombrado: {ren}")
    pol = next(e for e in despues["ej"] if e["name"] == "Elev laterales polea")
    check(pol["labels"][7] == "4 @ RIR1/0", "prescripción S8")
    check(despues["carrera"][5] == "3×(7'/1')", f"nivel de carrera: {despues['carrera'][5]}")
    check(len(despues["reglas"]) == len(antes["reglas"]) + 1 and despues["reglas"][-1] == "Regla de prueba [LÓGICA].",
          "regla nueva")
    for w in range(8):
        abs_ = despues["series"][w]["Abs"] - antes["series"][w]["Abs"]
        lat = despues["series"][w]["Hombro lat."] - antes["series"][w]["Hombro lat."]
        check(abs_ == (2 if w else 0), f"Abs S{w + 1}: +{abs_}")
        check(lat == (1 if w == 7 else 0), f"Hombro lat. S{w + 1}: +{lat}")

    # El Excel sigue bien: combinadas desplazadas y fórmulas de la tabla con la fila nueva.
    wb = openpyxl.load_workbook(COPIA)
    ws = wb[VOL]
    a = fila(ws, "Autorregulación")
    check(any(m.min_row == a and m.max_row == a for m in ws.merged_cells.ranges), "Autorregulación sigue combinada")
    abs_fila = next(r for r in range(1, ws.max_row + 1) if ws[f"A{r}"].value == "Abs" and r > a)
    check(f"LEFT(C{a - 1}," in ws[f"C{abs_fila}"].value, "la tabla de Abs suma la fila nueva")

    # Con el motor de fórmulas (si está): la tabla del Excel da lo mismo que la app.
    py = ROOT / "tools" / "out" / "py"
    probe = subprocess.run([sys.executable, "-c", "import formulas"], env={"PYTHONPATH": str(py), "SYSTEMROOT": "C:\\Windows"},
                           capture_output=True)
    if probe.returncode == 0:
        sys.path.insert(0, str(py))
        import formulas
        sol = formulas.ExcelModel().loads(str(COPIA)).finish().calculate()
        total = next(r for r in range(1, ws.max_row + 1) if str(ws[f"A{r}"].value or "").startswith("TOTAL"))
        for w, col in enumerate("CDEFGHIJ"):
            ref = f"'[{COPIA.name}]{VOL.upper()}'!{col}{total}"
            v = next((x.value[0][0] for k, x in sol.items() if k.upper() == ref.upper()), None)
            check(v == despues["series"][w]["TOTAL"], f"TOTAL S{w + 1}: Excel {v} ≠ app {despues['series'][w]['TOTAL']}")
    else:
        print("  --   sin el paquete «formulas» en tools/out/py: no se evalúa la tabla del Excel")

    # Una escalera nueva solo desde el Excel: extensión con rango que se amplía el viernes (pierna 2).
    esc = ROOT / "tools" / "out" / "prueba_escalera.xlsx"
    shutil.copyfile(XLSX, esc)
    wb = openpyxl.load_workbook(esc, rich_text=True)
    ws = wb[VOL]
    abd = next(r for r in range(1, ws.max_row + 1) if str(ws[f"A{r}"].value or "").startswith("Abducción de pie")
               and r > fila(ws, "(V)"))
    insertar_filas(ws, abd + 1, 1, estilo_de=abd)
    rangos = ["—", "2 x10-15/pierna\n90→45°", "2 x10-15/pierna\nnov: 90→35°", "2 x10-15/pierna\nnov: 90→25°",
              "2 x10-15/pierna\n90→25°", "2 x10-15/pierna\nnov: 90→15°", "2 x10-15/pierna\nnov: 90→5°",
              "2 x10-15/pierna\nnov: completo"]
    for col, v in zip("ABCDEFGHIJKL", ["Extensión unilateral (rango)", "Rehab"] + rangos
                      + ["Rango nuevo el viernes.", "id=extension-rango;escalera=extension"]):
        ws[f"{col}{abd + 1}"] = v
    wb.save(esc)
    regenerar(esc)
    motor = subprocess.run(["node", "-e", MOTOR_JS, str(esc)], cwd=ROOT, capture_output=True, encoding="utf-8")
    if motor.returncode:
        raise SystemExit(motor.stderr)
    m = json.loads(motor.stdout)
    check(m["escalera"]["max"] == [None, None, 0, 1, 2, 2, 3, 4, 5], f"escalera: {m['escalera']}")
    check(m["escalera"]["niveles"] == ["90→45°", "90→35°", "90→25°", "90→15°", "90→5°", "completo"], f"niveles: {m['escalera']['niveles']}")
    check(m["niveles"] == [0, 1, 2], f"niveles S2-S4 con todo en verde: {m['niveles']}")
    check(m["S4"] == "2 x10-15/pierna\n90→25° · nov", f"prescripción S4: {m['S4']!r}")
    check(m["S4retro"] == "2 x10-15/pierna\n90→35° · mantiene", f"con el viernes de S3 en Mantener: {m['S4retro']!r}")

    # Insertar y quitar la misma fila deja el plan y las celdas combinadas como estaban.
    ida = ROOT / "tools" / "out" / "prueba_ida_vuelta.xlsx"
    shutil.copyfile(XLSX, ida)
    wb = openpyxl.load_workbook(ida, rich_text=True)
    ws = wb[VOL]
    combinadas = sorted(str(m) for m in ws.merged_cells.ranges)
    a = fila(ws, "Autorregulación")
    insertar_filas(ws, a, 2)
    quitar_filas(ws, a, 2)
    check(sorted(str(m) for m in ws.merged_cells.ranges) == combinadas, "ida y vuelta: celdas combinadas")
    wb.save(ida)
    regenerar(ida)
    check(plan_de(ida) == antes, "ida y vuelta: el plan queda igual")

    for e in errores:
        print("FALLO", e)
    print(f"Edición simulada: {'todo OK' if not errores else f'{len(errores)} fallo(s)'}")
    return 1 if errores else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
