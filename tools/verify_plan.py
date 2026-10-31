"""Comprueba, con openpyxl, que el plan que construye la app (js/plan.js a partir del Excel
empotrado) es el de la hoja «Volumen (1910-1112)» del Excel de la raíz, y que el Excel
sirve para el iPhone.

    python tools/verify_plan.py [ruta.xlsx]

Genérico: localiza sesiones, ejercicios y textos por las etiquetas de la columna «App» (L)
— no por números de fila —, así que vale tras cualquier edición del plan. Comprueba:
  - sesiones (nombre y título), ejercicios (nombre, objetivo, id, celdas C–J, notas K) y
    el bloque de cardio y carrera;
  - textos: reglas, notas, hitos, títulos y reglas del bloque;
  - la tabla de series por grupo: cada celda es una fórmula normal que suma exactamente las
    filas de su grupo, y la app da el mismo recuento;
  - que no hay fórmulas matriciales en todo el libro (Excel para iPhone las calcula a 0);
  - que el Excel empotrado en la app (js/original-xlsx.js) es este mismo fichero.
Sale con código 1 si algo difiere.
"""
import base64
import json
import pathlib
import re
import subprocess
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
XLSX = ROOT / "Plan_8sem_vol_definitivo_con_cardio.xlsx"
SHEET = "Volumen (1910-1112)"
WEEKCOLS = "CDEFGHIJ"

# La app se carga en Node igual que en el navegador: el plan sale del Excel empotrado.
# (module/exports a undefined: así fflate se define como global, como en el navegador).
DUMP = r"""
const fs = require('fs');
const src = ['vendor/fflate.js', 'js/original-xlsx.js', 'js/lector-xlsx.js', 'js/plan.js']
  .map((f) => fs.readFileSync(f, 'utf8')).join('\n;\n');
const r = new Function('module', 'exports', src + ';return { PLAN, ORIGINAL_XLSX_B64 };')(undefined, undefined);
const PLAN = r.PLAN;
process.stdout.write(JSON.stringify({
  blocks: PLAN.BLOCK_VOL, rules: PLAN.VOL_RULES, notes: PLAN.VOL_NOTES,
  hitos: PLAN.VOL_HITOS, hitosNote: PLAN.VOL_HITOS_NOTE,
  blockRules: PLAN.VOL_BLOCK_RULES, blockRulesTitle: PLAN.VOL_BLOCK_RULES_TITLE,
  cardioTitle: PLAN.VOL_CARDIO_TITLE, grupos: PLAN.VOL_GRUPOS,
  series: [1, 2, 3, 4, 5, 6, 7, 8].map((n) => PLAN.seriesPorGrupo(n)),
  empotrado: r.ORIGINAL_XLSX_B64 }));
"""


def raw(v):
    """Texto de la celda tal cual (sin strip): los saltos de línea y espacios cuentan."""
    return "" if v is None else str(v)


def etiquetas(valor):
    out = {}
    for parte in raw(valor).split(";"):
        parte = parte.strip()
        if parte:
            k, _, v = parte.partition("=")
            out[k.strip()] = v.strip() if _ else True
    return out


def main():
    xlsx = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else XLSX
    out = subprocess.run(["node", "-e", DUMP], cwd=ROOT, capture_output=True, check=True, encoding="utf-8").stdout
    plan = json.loads(out)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[SHEET]
    app = lambda r: etiquetas(ws[f"L{r}"].value)

    errors = []
    compared = 0

    def check(ref, expected, actual):
        nonlocal compared
        compared += 1
        if expected != actual:
            errors.append(f"{ref}: Excel {expected!r} != app {actual!r}")

    def check_ejercicio(r, ex, nombre, objetivo):
        check(f"A{r} (nombre)", nombre, ex["name"])
        check(f"B{r}", objetivo, ex["target"])
        check(f"L{r} (id)", app(r).get("id", ex["id"]), ex["id"])
        for w, col in enumerate(WEEKCOLS):
            check(f"{col}{r}", raw(ws[f"{col}{r}"].value), ex["pres"][w]["label"])
        check(f"K{r}", raw(ws[f"K{r}"].value), ex["notes"])

    # Sesiones y sus ejercicios (filas con «id=»), y el cardio al final de su sesión.
    vistos = {}
    sesion = None
    for r in range(1, ws.max_row + 1):
        e = app(r)
        a = raw(ws[f"A{r}"].value)
        if e.get("sesion") is True:
            m = re.match(r"^\(([LMXJVSD])\)\s*([^·]+?)\s*(?:·\s*([\s\S]*))?$", a)
            sesion = m.group(1)
            s = plan["blocks"].get(sesion)
            if not s:
                errors.append(f"A{r}: la sesión {sesion} no está en la app")
                continue
            if sesion in vistos:   # otra cabecera del mismo día: sus filas siguen a las de la primera
                continue
            check(f"A{r} (día)", m.group(2).strip(), s["name"])
            check(f"A{r} (título)", (m.group(3) or "").strip(), s["title"])
            vistos[sesion] = 0
            continue
        if not a:
            sesion = None
        if "id" in e and not isinstance(e.get("sesion"), str) and sesion:
            ex = plan["blocks"][sesion]["exercises"][vistos[sesion]]
            vistos[sesion] += 1
            check_ejercicio(r, ex, a.strip(), raw(ws[f"B{r}"].value))
        elif isinstance(e.get("sesion"), str):
            letra = e["sesion"]
            vistos.setdefault(letra, 0)
            ex = plan["blocks"][letra]["exercises"][vistos[letra]]
            vistos[letra] += 1
            check_ejercicio(r, ex, raw(ws[f"B{r}"].value).strip(), "Cardio")
    for letra, n in vistos.items():
        if n != len(plan["blocks"][letra]["exercises"]):
            errors.append(f"Sesión {letra}: {n} ejercicios en el Excel, {len(plan['blocks'][letra]['exercises'])} en la app")

    # Textos por etiqueta.
    filas = lambda clave: [r for r in range(1, ws.max_row + 1) if clave in app(r)]
    check("reglas", [raw(ws[f"A{r}"].value) for r in filas("regla")], plan["rules"])
    check("notas", [raw(ws[f"A{r}"].value) for r in filas("nota")], plan["notes"])
    h = filas("hitos")[0]
    check("hitos", [raw(ws[f"{c}{h}"].value) for c in WEEKCOLS], plan["hitos"])
    check("nota de hitos", raw(ws[f"K{h}"].value), plan["hitosNote"])
    check("título del cardio", raw(ws[f"A{filas('cardio')[0]}"].value), plan["cardioTitle"])
    rb = filas("reglas-bloque")[0]
    check("título de las reglas del bloque", raw(ws[f"A{rb}"].value), plan["blockRulesTitle"])
    reglas = []
    r = rb + 1
    while ws[f"A{r}"].value:
        reglas.append(raw(ws[f"A{r}"].value))
        r += 1
    check("reglas del bloque", reglas, plan["blockRules"])

    # Series por grupo: fórmulas normales, que suman exactamente las filas de su grupo.
    g = filas("grupos")[0]
    r = g + 1
    while not raw(ws[f"A{r}"].value).startswith("TOTAL"):
        grupo = raw(ws[f"A{r}"].value)
        del_grupo = [f for f in range(1, g) if raw(ws[f"B{f}"].value) == grupo]
        for w, col in enumerate(WEEKCOLS):
            formula = ws[f"{col}{r}"].value
            if not isinstance(formula, str) or not formula.startswith("="):
                errors.append(f"{col}{r}: el recuento de {grupo!r} no es una fórmula normal")
                continue
            refs = sorted(int(x) for x in re.findall(rf"LEFT\({col}(\d+),", formula))
            if refs != del_grupo:
                errors.append(f"{col}{r}: suma las filas {refs}; las de {grupo!r} son {del_grupo}")
            propio = sum(int(raw(ws[f"{col}{f}"].value).split(" ")[0]) for f in del_grupo
                         if raw(ws[f"{col}{f}"].value).split(" ")[0].isdigit())
            check(f"{col}{r} ({grupo} S{w + 1})", propio, plan["series"][w][grupo])
        r += 1

    # Ninguna fórmula matricial en todo el libro.
    for hoja in wb:
        for fila in hoja.iter_rows():
            for celda in fila:
                if type(celda.value).__name__ == "ArrayFormula":
                    errors.append(f"{hoja.title}!{celda.coordinate}: fórmula matricial")

    # El Excel empotrado es este fichero.
    if base64.b64decode(plan["empotrado"]) != xlsx.read_bytes():
        errors.append("js/original-xlsx.js no es el Excel de la raíz: ejecuta python tools/excel_app.py")

    for e in errors:
        print("DIFERENCIA", e)
    print(f"{compared} celdas comparadas · {len(errors)} diferencias")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
