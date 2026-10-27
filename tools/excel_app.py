"""El Excel de la raíz es la fuente de verdad del plan: la app lo lleva empotrado y lo lee
al arrancar (js/lector-xlsx.js + js/plan.js + js/rodilla-data.js). Para cambiar el plan
basta con editar el Excel y ejecutar:

    python tools/excel_app.py

que hace, en orden:
  1. Regenera TODAS las fórmulas de «Volumen (1910-1112)» y «Seguimiento rodilla» a partir
     de las etiquetas de la columna «App» (oculta): así insertar, quitar o mover filas no
     deja referencias rotas (openpyxl no las actualiza al insertar filas).
  2. Comprueba que no queda ninguna fórmula matricial (Excel para iPhone las calcula a 0).
  3. Guarda el libro marcado para recalcular al abrirse.
  4. Vuelca el Excel a docs/plan-rodilla/ (tools/volcado_excel.py), lo empotra en la app
     (tools/embed.py) y sube la versión de la caché del service worker.

Si solo se ha tocado código (no el Excel): `python tools/excel_app.py --solo-cache`.

Etiquetas de la columna «App» (L en Volumen, R en Seguimiento): claves separadas por «;»,
con valor opcional («clave=valor»). Ver README, «Cambiar el plan».
"""
import pathlib
import re
import subprocess
import sys

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
XLSX = ROOT / "Plan_8sem_vol_definitivo_con_cardio.xlsx"
VOL, SEG = "Volumen (1910-1112)", "Seguimiento rodilla"
COL_APP = {VOL: "L", SEG: "R"}
SEMANAS = "CDEFGHIJ"   # S1..S8 en Volumen
CARRERA = "Carrera (≤2/10)"


def etiquetas(valor):
    """'id=prensa;escalera=prensa' -> {'id': 'prensa', 'escalera': 'prensa'}."""
    out = {}
    for parte in str(valor or "").split(";"):
        parte = parte.strip()
        if parte:
            k, _, v = parte.partition("=")
            out[k.strip()] = v.strip() if _ else True
    return out


def filas_con(ws, clave, valor=None):
    col = COL_APP[ws.title]
    return [r for r in range(1, ws.max_row + 1)
            if clave in (e := etiquetas(ws[f"{col}{r}"].value)) and (valor is None or e[clave] == valor)]


def una(ws, clave, valor=None):
    filas = filas_con(ws, clave, valor)
    if len(filas) != 1:
        raise SystemExit(f"{ws.title}: se esperaba una fila con «{clave}{'=' + valor if valor else ''}», hay {filas}")
    return filas[0]


def bloque(ws, cabecera, vale=lambda a: a not in (None, "")):
    """Filas de datos bajo una fila de cabecera, mientras la columna A cumpla `vale`
    (por defecto, hasta la primera vacía)."""
    filas = []
    r = cabecera + 1
    while vale(ws[f"A{r}"].value):
        filas.append(r)
        r += 1
    return filas


es_semana = lambda a: isinstance(a, str) and re.fullmatch(r"S\d+", a) is not None


def slug(texto):
    import unicodedata
    s = unicodedata.normalize("NFD", texto.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def entero_inicial(col, fila):
    return f'IFERROR(VALUE(LEFT({col}{fila},FIND(" ",{col}{fila}&" ")-1)),0)'


def formulas_volumen(ws):
    # Ejercicios (dentro de las tablas de sesión, antes de la tabla de grupos): id -> fila.
    grupos = una(ws, "grupos")
    ids = {}
    for r in range(1, grupos):
        a, b = ws[f"A{r}"].value, ws[f"B{r}"].value
        if a and b and a != "Ejercicio" and not str(a).startswith("("):
            e = etiquetas(ws[f"L{r}"].value)
            ids[e.get("id") or slug(str(a).strip())] = r

    # Series por grupo: suma, fila a fila, el entero inicial de los ejercicios de ese grupo.
    filas_grupo = bloque(ws, grupos)
    total = filas_grupo[-1]
    assert str(ws[f"A{total}"].value).startswith("TOTAL"), "la tabla de grupos debe acabar en TOTAL"
    for r in filas_grupo[:-1]:
        grupo = ws[f"A{r}"].value
        ejercicios = [f for f in range(1, grupos) if ws[f"B{f}"].value == grupo]
        if not ejercicios:
            raise SystemExit(f"{VOL}: ningún ejercicio con objetivo {grupo!r}")
        for col in SEMANAS:
            ws[f"{col}{r}"] = "=" + "+".join(entero_inicial(col, f) for f in ejercicios)
        ws[f"B{r}"] = f'=C{r}&" → "&J{r}'
    for col in SEMANAS:
        ws[f"{col}{total}"] = f"=SUM({col}{filas_grupo[0]}:{col}{filas_grupo[-2]})"
    ws[f"B{total}"] = f'=C{total}&" → "&J{total}'

    # Desglose de espalda por planos: «desglose=id1,id2…».
    for r in filas_con(ws, "desglose"):
        filas = []
        for i in etiquetas(ws[f"L{r}"].value)["desglose"].split(","):
            if i not in ids:
                raise SystemExit(f"{VOL}!L{r}: desglose con un id que no existe: {i!r}")
            filas.append(ids[i])
        for col in SEMANAS:
            ws[f"{col}{r}"] = "=" + "+".join(entero_inicial(col, f) for f in filas)
        ws[f"B{r}"] = f'=C{r}&" → "&J{r}'


def semaforo(r):
    """Columna P: semáforo de la sesión (reglas A104/A105 de la hoja Volumen)."""
    carrera = f'$E{r}="{CARRERA}"'
    falla_sesion = f'OR($J{r}>IF({carrera},2,3),AND({carrera},$K{r}="No"))'
    return (
        f'=IF($E{r}="—","—",IF($J{r}="","",IF($I{r}="","sin referencia",'
        f'IF(AND({carrera},$K{r}=""),"falta K",IF({falla_sesion},"Retroceder",'
        f'IF($M{r}="","pendiente",IF($M{r}>$I{r},"Retroceder",'
        f'IF($N{r}="","pendiente",IF($O{r}<=$I{r},"Avanzar","Mantener")))))))))')


def formulas_seguimiento(ws):
    dias = bloque(ws, una(ws, "dias"))
    d0, d1 = dias[0], dias[-1]
    c_, i_ = f"$C${d0}:$C${d1}", f"$I${d0}:$I${d1}"
    for k, r in enumerate(dias):
        ws[f"I{r}"] = f'=IF(COUNT(G{r},H{r})=0,"",MAX(G{r},H{r}))'
        # 24 h / 48 h: las mañanas de las filas siguientes (tras la última no hay datos).
        ws[f"M{r}"] = f'=IF(I{r + 1}="","",I{r + 1})' if k + 1 < len(dias) else '=""'
        ws[f"N{r}"] = f'=IF(I{r + 2}="","",I{r + 2})' if k + 2 < len(dias) else '=""'
        ws[f"O{r}"] = f'=IF(OR(M{r}="",N{r}=""),"",MAX(M{r},N{r}))'
        ws[f"P{r}"] = semaforo(r)

    # Resumen semanal (sin fórmulas matriciales: cada semana son filas seguidas).
    filas_sem = bloque(ws, una(ws, "semanas"), es_semana)
    fila_de = {ws[f"A{r}"].value: r for r in filas_sem}
    for r in filas_sem:
        s = ws[f"A{r}"].value
        propias = [x for x in dias if ws[f"C{x}"].value == s]
        a, b = propias[0], propias[-1]
        assert propias == list(range(a, b + 1)), f"los días de {s} no son filas seguidas"
        ws[f"B{r}"] = (f'=IF(COUNTIFS({c_},$A{r},{i_},">=0")=0,"",ROUND(SUMIFS({i_},{c_},$A{r})'
                       f'/COUNTIFS({c_},$A{r},{i_},">=0"),1))')
        ws[f"E{r}"] = f'=IF(COUNT(I{a}:I{b})=0,"",MAX(I{a}:I{b}))'
        ws[f"F{r}"] = f'=COUNTIFS({c_},$A{r},$P${d0}:$P${d1},"Retroceder")'
        ws[f"G{r}"] = f'=COUNTIFS({c_},$A{r},$L${d0}:$L${d1},"Sí")'

    # Avisos: S2 (semana de referencia), S4 (fisio) y chasquido.
    ref = ws[f"A{filas_sem[0]}"].value
    de_ref = [x for x in dias if ws[f"C{x}"].value == ref]
    primera = '"x"'
    for x in reversed(de_ref):
        primera = f"IF(ISNUMBER(I{x}),I{x},{primera})"
    rango_ref = f"I{de_ref[0]}:I{de_ref[-1]}"
    ws[f"E{una(ws, 'aviso', 's2')}"] = (
        f'=IF(COUNT({rango_ref})=0,"Pendiente: falta la primera medición de {ref}",'
        f'IF(MAX({rango_ref})>{primera},'
        f'"REPITE {ref}: alguna mañana salió peor que la referencia; en S{int(ref[1:]) + 1} no sube nada",'
        f'"{ref} estable por ahora: en S{int(ref[1:]) + 1} se sigue el plan"))')
    s4 = [x for x in dias if ws[f"C{x}"].value == "S4"]
    cierre = ws[f"A{s4[-1]}"].value
    ddmm = f"{cierre.day}/{cierre.month}"
    b2, b4 = f"B{fila_de[ref]}", f"B{fila_de['S4']}"
    ws[f"E{una(ws, 'aviso', 's4')}"] = (
        f'=IF(TODAY()<DATE({cierre.year},{cierre.month},{cierre.day}),"Pendiente: se calcula al cerrar S4 (domingo {ddmm})",'
        f'IF(OR(COUNTIFS({c_},"{ref}",{i_},">=0")<5,COUNTIFS({c_},"S4",{i_},">=0")<5),'
        f'"Pendiente: hacen falta al menos 5 mañanas registradas en {ref} y en S4",'
        f'IF(ROUND({b4}-{b2},1)>=0.5,"PEOR que en {ref} (+0,5 o más): para la progresión y la valoración del fisio pasa a ser prioritaria",'
        f'IF(ROUND({b2}-{b4},1)>=0.5,"MEJOR que en {ref} (−0,5 o más): sigue el plan",'
        f'"IGUAL que en {ref} (menos de 0,5 de diferencia): la valoración del fisio pasa a ser prioritaria"))))')
    ws[f"E{una(ws, 'aviso', 'chasquido')}"] = (
        f'=IF(COUNTIF($L${d0}:$L${d1},"Sí")>0,"Registrado: la valoración del fisio pasa a ser prioritaria",'
        f'"Sin chasquido doloroso por ahora")')

    # Entrada a la carrera.
    walk, td, hops = una(ws, "entrada", "walk"), una(ws, "entrada", "touchdown"), una(ws, "entrada", "hops")
    manana = f"INDEX({i_},MATCH({{d}},$A${d0}:$A${d1},0))"
    hoy, sig = manana.format(d=f"G{walk}"), manana.format(d=f"G{walk}+1")
    ws[f"F{walk}"] = f'=IF(G{walk}="","",IFERROR(IF(OR({hoy}="",{sig}=""),"",IF({sig}<={hoy},"Sí","No")),""))'
    for r in (td, hops):
        ws[f"G{r}"] = f"=G{walk}+1"
        ws[f"H{r}"] = f'=IF(D{r}="","",D{r})'
    ws[f"H{walk}"] = (f'=IF(OR(D{walk}="No",E{walk}="No",F{walk}="No"),"No",'
                      f'IF(AND(D{walk}="Sí",E{walk}="Sí",F{walk}="Sí"),"Sí",""))')
    res = una(ws, "entrada-resultado")
    rango = f"H{min(walk, td, hops)}:H{max(walk, td, hops)}"
    ws[f"H{res}"] = (f'=IF(COUNTIF({rango},"Sí")=3,"SÍ: 1.ª carrera hoy (N1)",IF(COUNTIF({rango},"No")>0,'
                     f'"NO: esta semana no hay carrera (el domingo, como mucho, un paseo); repite los criterios el miércoles siguiente","Pendiente"))')

    # Test de fuerza: estimada (Epley), simetría D/I y ≥90 %.
    for r in filas_con(ws, "test"):
        ws[f"I{r}"] = f'=IF(OR(F{r}="",G{r}=""),"",ROUND(F{r}*(1+G{r}/30),1))'
        ws[f"J{r}"] = f'=IF(OR(F{r}="",H{r}=""),"",ROUND(F{r}*(1+H{r}/30),1))'
        ws[f"K{r}"] = f'=IF(OR(I{r}="",J{r}=""),"",I{r}/J{r})'
        ws[f"L{r}"] = f'=IF(K{r}="","",IF(K{r}>=0.9,"Sí","No"))'

    # Rangos de validación y formato condicional ligados a la tabla de días.
    rango = lambda col: f"{col}{d0}:{col}{d1}"
    for dv in ws.data_validations.dataValidation:
        if dv.type == "whole":
            dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(f"{rango('G')} {rango('H')} {rango('J')}")
        elif dv.type == "list" and "K" in str(dv.sqref):
            dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(f"{rango('K')} {rango('L')}")
    cf = ws.conditional_formatting
    viejas = {str(k.sqref)[0]: v for k, v in list(cf._cf_rules.items()) if str(k.sqref)[0] in "PLJ" and ":" in str(k.sqref)}
    for k in [k for k in list(cf._cf_rules) if str(k.sqref)[0] in "PLJ" and ":" in str(k.sqref)]:
        del cf._cf_rules[k]
    for col, reglas in viejas.items():
        for regla in reglas:
            f = [re.sub(r"([A-Z]\$?)\d+", lambda m: m.group(1) + str(d0), x) for x in regla.formula]
            cf.add(rango(col), FormulaRule(formula=f, font=regla.dxf.font if regla.dxf else None,
                                           fill=regla.dxf.fill if regla.dxf else None))


def _mover(desde, n):
    """Función que desplaza un rango tras insertar (n > 0) o quitar (n < 0) filas en `desde`."""
    from openpyxl.worksheet.cell_range import CellRange

    def mover(rango):
        cr = CellRange(str(rango))
        if cr.min_row >= desde:
            cr.shift(row_shift=n)
        elif cr.max_row >= desde:
            cr.expand(down=n)
        return cr
    return mover


def _cambiar_filas(ws, desde, n, operar):
    """Inserta o quita filas (`operar`) y mueve lo que openpyxl no mueve: celdas combinadas,
    alturas de fila, validaciones y formato condicional. Las combinadas se deshacen antes y
    se rehacen después, desplazadas (openpyxl ya ha movido sus celdas)."""
    from openpyxl.worksheet.cell_range import MultiCellRange
    mover = _mover(desde, n)
    combinadas = [str(m) for m in ws.merged_cells.ranges]
    for c in combinadas:
        ws.unmerge_cells(c)
    alturas = {r: d.height for r, d in ws.row_dimensions.items() if d.height}
    operar()
    for c in combinadas:
        cr = mover(c)
        if cr.max_row >= cr.min_row and cr.min_row >= 1 and (cr.max_row > cr.min_row or cr.max_col > cr.min_col):
            ws.merge_cells(cr.coord)
    for r in list(ws.row_dimensions):
        ws.row_dimensions[r].height = None
    for r, h in alturas.items():
        nueva = r + n if r >= desde else r
        if nueva >= 1:
            ws.row_dimensions[nueva].height = h
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(" ".join(mover(r).coord for r in dv.sqref.ranges))
    cf = ws.conditional_formatting
    for rango, reglas in list(cf._cf_rules.items()):
        nuevo = " ".join(mover(r).coord for r in rango.sqref.ranges)
        del cf._cf_rules[rango]
        for regla in reglas:
            regla.formula = [re.sub(r"(\$?[A-Z]{1,2}\$?)(\d+)",
                                    lambda m: m.group(1) + str(int(m.group(2)) + n if int(m.group(2)) >= desde else int(m.group(2))), f)
                             for f in regla.formula]
            cf.add(nuevo, regla)


def insertar_filas(ws, antes_de, n=1, estilo_de=None):
    """Inserta n filas antes de la fila `antes_de`, con el estilo de la fila `estilo_de`
    (numerada antes de insertar; por defecto, la de encima). Las fórmulas las rehace
    `regenerar()` a partir de las etiquetas."""
    estilo_de = estilo_de or antes_de - 1
    _cambiar_filas(ws, antes_de, n, lambda: ws.insert_rows(antes_de, n))
    origen = estilo_de + n if estilo_de >= antes_de else estilo_de
    for r in range(antes_de, antes_de + n):
        ws.row_dimensions[r].height = ws.row_dimensions[origen].height
        for col in range(1, ws.max_column + 1):
            ws.cell(r, col)._style = ws.cell(origen, col)._style


def quitar_filas(ws, desde, n=1):
    """Quita n filas desde la fila `desde` (incluida)."""
    _cambiar_filas(ws, desde + n, -n, lambda: ws.delete_rows(desde, n))


def sin_matrices(wb):
    malas = [f"{ws.title}!{c.coordinate}" for ws in wb for fila in ws.iter_rows() for c in fila
             if type(c.value).__name__ == "ArrayFormula"]
    if malas:
        raise SystemExit(f"Fórmulas matriciales (Excel para iPhone las calcula a 0): {malas[:10]}…")


def regenerar(path=XLSX):
    wb = openpyxl.load_workbook(path, rich_text=True)
    formulas_volumen(wb[VOL])
    formulas_seguimiento(wb[SEG])
    sin_matrices(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def version_cache():
    """CACHE del service worker = huella de todos los ficheros que cachea."""
    import hashlib
    sw = ROOT / "sw.js"
    texto = sw.read_text(encoding="utf-8")
    assets = re.findall(r"'\./([^']*)'", texto.split("const ASSETS")[1].split("];")[0])
    h = hashlib.sha1()
    for a in assets:
        if a:
            h.update((ROOT / a).read_bytes())
    nuevo = re.sub(r"const CACHE = '[^']*';", f"const CACHE = 'gym-{h.hexdigest()[:8]}';", texto)
    sw.write_text(nuevo, encoding="utf-8", newline="\n")
    return nuevo.split("const CACHE = '")[1].split("'")[0]


def main():
    if "--solo-cache" in sys.argv:   # se tocó código pero no el Excel
        print("sw.js: CACHE =", version_cache())
        return
    regenerar()
    print(f"{XLSX.name}: fórmulas regeneradas, sin matrices")
    for script in ("volcado_excel.py", "embed.py"):
        subprocess.run([sys.executable, str(ROOT / "tools" / script)], check=True)
    print("sw.js: CACHE =", version_cache())


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
