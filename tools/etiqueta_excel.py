"""Histórico (27/10/2026): añadió al Excel la columna «App» con las etiquetas que la app
lee para construir el plan. Ya está aplicado; si lo vuelves a ejecutar, reescribe las
mismas etiquetas. Las siguientes ediciones del plan se hacen en el Excel y se publican
con `python tools/excel_app.py`.

    python tools/etiqueta_excel.py
"""
import pathlib
import sys

import openpyxl
from openpyxl.styles import Font

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from excel_app import SEG, VOL, XLSX, slug  # noqa: E402

GRIS = Font(name="Arial", size=8, color="FF8A939C")


def fila(ws, texto, col="A", empieza=False):
    filas = [r for r in range(1, ws.max_row + 1)
             if isinstance(ws[f"{col}{r}"].value, str)
             and (ws[f"{col}{r}"].value.startswith(texto) if empieza else ws[f"{col}{r}"].value.strip() == texto)]
    assert len(filas) == 1, (texto, filas)
    return filas[0]


def poner(ws, col, r, valor):
    ws[f"{col}{r}"] = valor
    ws[f"{col}{r}"].font = GRIS


def volumen(ws):
    # Cabeceras de sesión y de tabla.
    for r in range(1, ws.max_row + 1):
        a = ws[f"A{r}"].value
        if isinstance(a, str) and a.startswith("(") and ")" in a[:4]:
            poner(ws, "L", r, "sesion;nombreS1=Sábado/Domingo" if a.startswith("(S)") else "sesion")
        elif a in ("Ejercicio", "Día / Sesión"):
            poner(ws, "L", r, "App (no tocar)")
    # Ejercicios: id estable (el slug de su nombre de hoy) y su papel en la progresión.
    papel = {
        "Prensa unil · pesado y lento n.º 1": "escalera=prensa",
        "Sentadilla al cajón o desde pines · pesado y lento n.º 2": "escalera=prensa:sigue",
        "Bulgarian split squat@J": "escalera=prensa:sigue",
        "Touchdown squat@L": "escalera=prensa:anterior",
    }
    sesion = None
    grupos = fila(ws, "Grupo muscular")
    for r in range(1, grupos):
        a, b = ws[f"A{r}"].value, ws[f"B{r}"].value
        if isinstance(a, str) and a.startswith("(") and ")" in a[:4]:
            sesion = a[1]
            continue
        if a and b and a != "Ejercicio":
            nombre = str(a).strip()
            extra = papel.get(nombre) or papel.get(f"{nombre}@{sesion}")
            poner(ws, "L", r, f"id={slug(nombre)}" + (f";{extra}" if extra else ""))
    # Bloque de cardio y carrera.
    poner(ws, "L", fila(ws, "BLOQUE DE CARDIO", empieza=True), "cardio")
    for texto, tags in (("Martes (Post-Push)", "sesion=M;bici"), ("Viernes (Post-Torso)", "sesion=V;bici"),
                        ("Miércoles · Carrera (nivel nuevo)", "sesion=X;escalera=carrera"),
                        ("Domingo · Carrera (repite)", "sesion=D;titulo=Carrera;escalera=carrera:domingo")):
        r = fila(ws, texto)
        poner(ws, "L", r, f"id={slug(ws[f'B{r}'].value.strip())};{tags}")
    poner(ws, "L", fila(ws, "Hitos y avisos"), "hitos")
    # Textos.
    poner(ws, "L", 2, "regla")
    poner(ws, "L", fila(ws, "Autorregulación:", empieza=True), "regla")
    for texto in ("Semana 9 (14-20/12)", "Hipertrofia: la semana sin entrenar", "Rodilla en el viaje:", "Semana 10 (vuelta):"):
        poner(ws, "L", fila(ws, texto, empieza=True), "nota")
    poner(ws, "L", grupos, "grupos")
    poner(ws, "L", fila(ws, "REGLAS DEL BLOQUE RODILLA", empieza=True), "reglas-bloque")
    # Desglose de espalda por planos (ids de los ejercicios que suma cada fila).
    for texto, ids in (
            ("Dorsal Bajo (Plano Frontal)", "jalon-ancho-plano-frontal,aduccion-unil-polea-alta-pull-in"),
            ("Dorsal Superior (Plano Sagital)", "jalon-neutro,dominada-neutra-dorsal,remo-maquina-pecho-apoyado-30"),
            ("Espalda Alta (Plano Transversal)", "remo-gironda-abierto"),
            ("Total Transversal (Espalda Alta + Post.)", "remo-gironda-abierto,posterior-polea")):
        poner(ws, "L", fila(ws, texto), f"desglose={ids}")
    ws.column_dimensions["L"].hidden = True


def seguimiento(ws):
    col = "R"
    poner(ws, col, fila(ws, "Fecha"), "dias")
    poner(ws, col, fila(ws, "Semana"), "semanas")
    poner(ws, col, fila(ws, "S2 · semana de referencia"), "aviso=s2")
    poner(ws, col, fila(ws, "AVISO S4 (fisio)"), "aviso=s4")
    poner(ws, col, fila(ws, "Chasquido doloroso"), "aviso=chasquido")
    poner(ws, col, fila(ws, "La molestia de cada semana", empieza=True), "texto=resumen")
    poner(ws, col, fila(ws, "Entrada a la carrera", empieza=True), "titulo=entrada")
    for texto, tags in (("Paseo de 30 min", "entrada=walk;dia=martes"),
                        ("10 touchdown con cada pierna", "entrada=touchdown;dia=miércoles"),
                        ("20 saltitos a dos piernas", "entrada=hops;dia=miércoles")):
        poner(ws, col, fila(ws, texto, empieza=True), tags)
    poner(ws, col, fila(ws, "Resultado"), "entrada-resultado")
    poner(ws, col, fila(ws, "Test de fuerza", empieza=True), "titulo=fuerza")
    for i, texto in enumerate(("Test 1 (S2)", "Test 2 (S5)", "Test 3 (S7)", "Opcional (S8)"), 1):
        poner(ws, col, fila(ws, texto), f"test=t{i}")
    poner(ws, col, fila(ws, "Fuerza estimada = carga", empieza=True), "texto=fuerza")
    poner(ws, col, fila(ws, "Escalones de profundidad", empieza=True), "titulo=escalones")
    for i, texto in enumerate(("0 (actual)", "1", "2", "3", "4 (= I)")):
        poner(ws, col, fila(ws, texto), f"escalon={i}")
    poner(ws, col, fila(ws, "Reparte la diferencia", empieza=True), "texto=escalones")
    ws.column_dimensions[col].hidden = True


def main():
    wb = openpyxl.load_workbook(XLSX, rich_text=True)
    volumen(wb[VOL])
    seguimiento(wb[SEG])
    wb.calculation.fullCalcOnLoad = True
    wb.save(XLSX)
    print("etiquetas añadidas a", XLSX.name)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
